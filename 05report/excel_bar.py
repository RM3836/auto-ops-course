from openpyxl import load_workbook
from openpyxl.chart import BarChart3D, Reference  # BarChart3D为三维柱形图
wb = load_workbook('servertest.xlsx')             # 打开xlsx文件
ws = wb.active                                         # 获取当前的工作表对象（第1个工作表）
# 选择图表的数据源,返回范围内的所有单元格
data = Reference(ws, min_col=3, min_row=2, max_col=3, max_row=8)
# 选择要显示的X轴坐标标记内容
x_label = Reference(ws, min_col= 2, min_row = 2, max_row = 8)
chart  = BarChart3D()                    # 创建BarChart3D对象（三维柱形图表）
chart.title = "服务器性能测试"      # 给BarChart3D对象添加标题
chart.add_data(data)               # 给BarChart3D对象添加数据源
chart.set_categories(x_label)    # 给BarChart3D对象添加坐标轴标记
ws.add_chart(chart, "E5")         # 在工作表上添加图表，并指定图表左上角锚定的单元格
wb.save("server_bar3d.xlsx")      # 保存工作薄
